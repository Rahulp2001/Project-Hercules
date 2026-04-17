"""Generate Project Hercules Planning Document v2.0 (Full-Stack) as a Word file."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)


def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    return h


def add_table_from_data(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), '1A1A2E')
        shading.set(qn('w:val'), 'clear')
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
            for p in row.cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_note(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0xCC, 0x88, 0x00)
    run.font.size = Pt(10)


def add_code_block(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)


def add_new_section_tag():
    p = doc.add_paragraph()
    run = p.add_run('NEW IN V2.0')
    run.bold = True
    run.font.color.rgb = RGBColor(0x00, 0xB8, 0x94)
    run.font.size = Pt(10)


# ============================================================
# TITLE PAGE
# ============================================================
for _ in range(6):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('PROJECT HERCULES')
run.bold = True
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Personal Health & Fitness Goal Tracker')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x00, 0xB8, 0x94)

subtitle2 = doc.add_paragraph()
subtitle2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle2.add_run('Planning & Requirements Document')
run.font.size = Pt(14)

for _ in range(4):
    doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('Author: Rahul Prabhu\n').font.size = Pt(12)
info.add_run('Date: April 4, 2026\n').font.size = Pt(12)
info.add_run('Version: 2.0 (Full-Stack)').font.size = Pt(12)

doc.add_page_break()

# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================
add_heading_styled('1. Executive Summary', level=1)
doc.add_paragraph(
    'Project Hercules is a full-stack web application for personal fitness and health goal tracking. '
    'It guides users through an onboarding questionnaire, calculates their maintenance calories (TDEE), '
    'suggests a nutrition plan (cut/bulk/maintain), and provides daily tracking tools for calories, macros, '
    'water, sleep, steps, workouts, and custom goals.'
)
doc.add_paragraph('The deliverables are:')
add_bullet('A full-stack web application (React frontend + Express backend + PostgreSQL database)')
add_bullet('This planning document')
doc.add_paragraph()
add_note(
    'Clarification: This is a personal-use fitness tracker, not a clinical or medical tool. '
    'All calculations use standard fitness formulas (Mifflin-St Jeor). No medical advice is provided.'
)

# ============================================================
# 2. TARGET USERS & PERSONAS
# ============================================================
add_heading_styled('2. Target Users & Personas', level=1)
add_table_from_data(
    ['Persona', 'Description', 'Primary Tasks'],
    [
        ['New User', 'First-time visitor, no profile', 'Complete onboarding questionnaire, set initial goals'],
        ['Daily User (Primary)', 'Returning user tracking fitness', 'Log meals, workouts, water, sleep; review dashboard'],
        ['Goal Reviewer', 'User checking weekly/monthly progress', 'View charts, analyze trends, adjust targets'],
        ['Power User', 'User who customizes everything', 'Create custom targets, adjust macros, modify routines'],
    ]
)
add_note('MVP: All personas are the same user at different stages. Single-user app (no multi-user auth for MVP).')

# ============================================================
# 3. MVP VS. FUTURE SCOPE
# ============================================================
add_heading_styled('3. MVP vs. Future Scope', level=1)

add_heading_styled('MVP (Phase 1-9)', level=2)
for item in [
    'Onboarding questionnaire (5 questions)',
    'TDEE calculation with cut/bulk/maintain suggestion',
    'Daily dashboard with all-in-one overview',
    'Calorie & macro tracker (food logging)',
    'Water intake tracker',
    'Sleep tracker',
    'Step counter (manual entry)',
    'Gym workout logger (exercises, sets, reps, weight)',
    'Running/cardio logger (distance, time, type)',
    'Custom daily targets (user-defined)',
    'Progress charts (weight, calories, weekly completion, macros)',
    'Streak tracking & milestone badges',
    'Best habit / worst habit analysis',
    'Daily motivational quotes',
    'Weekly completion summary',
    'Dark/light mode',
    'Responsive design (desktop + mobile browser)',
    'PostgreSQL database persistence',
    'REST API backend with full test coverage',
]:
    add_bullet(item)

add_heading_styled('Post-MVP (Future)', level=2)
for item in [
    'User authentication (JWT login/signup)',
    'Multi-user support with cloud sync',
    'Social sharing (share progress screenshots)',
    'Workout plan generator (AI-suggested routines)',
    'Food database with autocomplete (API integration)',
    'Barcode scanner for packaged food',
    'Body measurement tracker (waist, chest, arms, etc.)',
    'Progress photo journal',
    'Export data to CSV/PDF',
    'PWA (Progressive Web App) for mobile install',
    'Push notification reminders',
    'Integration with Google Fit / Apple Health APIs',
    'Android/iOS native app (React Native conversion)',
    'Deployment to Vercel + Railway with shareable URL',
]:
    add_bullet(item)

# ============================================================
# 4. SYSTEM WORKFLOW
# ============================================================
add_heading_styled('4. System Workflow', level=1)

add_heading_styled('4.1 First-Time User Flow (Onboarding)', level=2)
add_table_from_data(
    ['Step', 'Name', 'Description'],
    [
        ['1', 'Welcome Screen', 'Animated splash: "Welcome to Project Hercules!" with logo and tagline'],
        ['2', 'Name', '"What should we call you?" - Text input'],
        ['3', 'Body Stats', 'Age, Gender (Male/Female/Other), Weight (kg or lbs toggle), Height (cm or ft/in toggle)'],
        ['4', 'Activity Level', '5 options: Sedentary / Lightly Active / Moderately Active / Very Active / Extremely Active - with descriptions and examples'],
        ['5', 'Fitness Goal', '3 options: Lose Weight (Cut) / Build Muscle (Bulk) / Stay Healthy (Maintain) - with visual icons'],
        ['6', 'Results Screen', 'Shows calculated TDEE, suggested daily calories, macro split, and a motivational message'],
        ['7', 'Dashboard', 'Redirect to main dashboard. Onboarding complete.'],
    ]
)

add_heading_styled('4.2 TDEE Calculation (Mifflin-St Jeor Equation)', level=2)
add_table_from_data(
    ['Detail', 'Specification'],
    [
        ['Formula (Male)', 'BMR = (10 x weight_kg) + (6.25 x height_cm) - (5 x age) + 5'],
        ['Formula (Female)', 'BMR = (10 x weight_kg) + (6.25 x height_cm) - (5 x age) - 161'],
        ['Formula (Other)', 'Average of Male and Female formulas'],
        ['TDEE', 'BMR x Activity Multiplier'],
        ['Activity Multipliers', 'Sedentary: 1.2, Light: 1.375, Moderate: 1.55, Very: 1.725, Extreme: 1.9'],
        ['Cut Calories', 'TDEE - 500 (for ~0.5kg/week loss)'],
        ['Bulk Calories', 'TDEE + 300 (for lean bulk)'],
        ['Maintain Calories', 'TDEE (no adjustment)'],
    ]
)

add_heading_styled('4.3 Default Macro Split', level=2)
add_table_from_data(
    ['Goal', 'Protein', 'Carbs', 'Fats'],
    [
        ['Cut', '40%', '30%', '30%'],
        ['Bulk', '30%', '45%', '25%'],
        ['Maintain', '30%', '40%', '30%'],
    ]
)
doc.add_paragraph('User can customize these percentages in Settings. Must always total 100%.')

add_heading_styled('4.4 Daily Data Flow', level=2)
doc.add_paragraph('Morning:')
add_bullet('User opens app -> Dashboard loads with today\'s data from backend')
add_bullet('Previous night\'s sleep auto-prompted if not logged')
add_bullet('Daily motivational quote fetched from API')
doc.add_paragraph('Throughout Day:')
add_bullet('User logs meals -> POST to /api/meals -> dashboard updates in real-time')
add_bullet('User logs water -> POST to /api/water -> visual fill animation')
add_bullet('User logs workouts -> POST to /api/workouts -> exercise cards appear')
add_bullet('User logs steps -> POST to /api/steps -> step counter updates')
add_bullet('Custom targets checked off -> POST to /api/targets/daily')
doc.add_paragraph('End of Day:')
add_bullet('Daily completion percentage calculated server-side')
add_bullet('Streak updated (increment if >=80% targets met)')
add_bullet('Achievement badges checked and awarded if conditions met')
add_bullet('Weight entry prompted (optional, for trend tracking)')

add_heading_styled('4.5 Streak Logic', level=2)
add_table_from_data(
    ['Rule', 'Specification'],
    [
        ['Increment', 'User completes >=80% of their daily targets'],
        ['Targets counted', 'Calorie target (within +/-10%), water goal, workout logged, sleep logged, custom targets'],
        ['Reset', 'If a day is missed entirely (no data logged)'],
        ['Freeze', 'Future: "Rest day" option that preserves streak without logging'],
        ['Best Streak', 'Tracked separately, never resets'],
    ]
)

# ============================================================
# 5. FEATURE SPECIFICATIONS
# ============================================================
add_heading_styled('5. Feature Specifications', level=1)

add_heading_styled('5.1 Daily Dashboard (Home Screen)', level=2)
doc.add_paragraph('The dashboard is the central hub. All tracking data for today is visible at a glance.')
add_table_from_data(
    ['Widget', 'Content', 'Interaction'],
    [
        ['Greeting Banner', '"Good morning, Rahul!" (time-aware) + motivational quote', 'Quote changes daily'],
        ['Calorie Ring', 'Circular progress: consumed / target with remaining in center', 'Tap -> opens food logger'],
        ['Macro Bars', '3 horizontal bars: Protein / Carbs / Fats with grams + percentage', 'Tap -> opens macro detail'],
        ['Water Tracker', 'Glass/bottle icons that fill up: "6/8 glasses"', 'Tap glass -> +1, long-press -> custom amount'],
        ['Sleep Card', 'Moon icon + "7.5 hrs" with quality badge (Good/Fair/Poor)', 'Tap -> edit sleep entry'],
        ['Steps Card', 'Shoe icon + "8,500 steps" with mini progress arc', 'Tap -> edit steps'],
        ['Workout Summary', 'Dumbbell icon + "2 workouts today" with exercise names', 'Tap -> opens workout log'],
        ['Custom Targets', 'Checklist of user-defined targets with checkboxes', 'Tap checkbox -> toggle'],
        ['Streak Counter', 'Fire icon + "5 day streak!" with flame animation at milestones', 'Tap -> shows streak history'],
        ['Daily Completion', 'Bottom progress bar: "Today: 75% complete"', 'Color: red <50%, yellow 50-79%, green >=80%'],
        ['Quick Actions FAB', 'Floating action button with + icon', 'Expands: Log Meal, Log Workout, Log Water, Log Weight'],
    ]
)

add_heading_styled('5.2 Calorie & Macro Tracker', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Meal Categories', 'Breakfast, Lunch, Dinner, Snacks'],
        ['Food Entry', 'Name (text), Calories (number), Protein (g), Carbs (g), Fats (g)'],
        ['Quick Add', 'Calorie-only quick entry (skip macros)'],
        ['Recent Foods', 'Last 20 logged foods, one-tap re-log'],
        ['Favorite Foods', 'User can star frequently eaten foods'],
        ['Meal Summary', 'Per-meal calorie subtotal + macro breakdown'],
        ['Daily Total', 'Running total across all meals vs. target'],
        ['Over/Under Indicator', 'Green = within +/-10%, Yellow = +/-10-20%, Red = >20% over/under'],
        ['Macro Visualization', 'Donut chart: protein (blue), carbs (orange), fats (yellow)'],
    ]
)

add_heading_styled('5.3 Water Intake Tracker', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Default Goal', '8 glasses (2L). Configurable in settings.'],
        ['Unit Options', 'Glasses (250ml each) or Liters'],
        ['Quick Add', 'Tap glass icon -> +1 glass. Each glass animates filling.'],
        ['Custom Amount', 'Long-press or button -> enter custom ml/oz'],
        ['Visual', 'Row of 8 glass icons. Filled = blue, empty = gray outline.'],
        ['Progress Text', '"1.5L / 2.0L" below the glasses'],
    ]
)

add_heading_styled('5.4 Sleep Tracker', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Input Method', 'Bedtime + Wake time -> auto-calculate hours, OR direct hours entry'],
        ['Quality Rating', '3 options: Good (green), Fair (yellow), Poor (red) - with emoji faces'],
        ['Ideal Range', '7-9 hours highlighted as green zone on weekly chart'],
        ['Weekly View', 'Bar chart showing sleep hours per day, colored by quality'],
        ['Average', '"This week: avg 7.2 hrs"'],
    ]
)

add_heading_styled('5.5 Step Counter', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Input', 'Manual number entry (no device sensor for MVP)'],
        ['Default Goal', '10,000 steps. Configurable.'],
        ['Visual', 'Circular progress arc with step count in center'],
        ['Weekly View', 'Bar chart of daily steps'],
    ]
)

add_heading_styled('5.6 Gym Workout Logger', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Workout Structure', 'Workout -> Exercises -> Sets'],
        ['Exercise Entry', 'Name, Category (dropdown), Notes'],
        ['Set Entry', 'Reps, Weight (kg/lbs), type (Normal / Warmup / Drop set / Failure)'],
        ['Categories', 'Chest, Back, Shoulders, Arms, Legs, Core, Full Body, Other'],
        ['Templates', 'Save a workout as template, one-tap re-use'],
        ['Recent Workouts', 'Last 10 workouts, quick duplicate'],
        ['Timer', 'Built-in rest timer between sets (configurable: 30s, 60s, 90s, 2min, 3min, custom)'],
        ['Volume Tracking', 'Total volume = sum of (reps x weight) per workout'],
        ['Personal Records', 'Auto-detect PR (heaviest weight or most reps) - highlighted with trophy icon'],
    ]
)

add_heading_styled('5.7 Running / Cardio Logger', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Activity Types', 'Running, Cycling, Swimming, Walking, Jump Rope, Elliptical, Other'],
        ['Fields', 'Duration (minutes), Distance (km/mi - optional), Calories burned (optional/calculated)'],
        ['Calorie Estimation', 'If not entered manually: MET x weight_kg x duration_hrs'],
        ['History', 'Card-based list of past cardio sessions'],
    ]
)
doc.add_paragraph('MET Values (for calorie estimation):')
add_table_from_data(
    ['Activity', 'MET'],
    [
        ['Running (moderate)', '9.8'],
        ['Cycling (moderate)', '7.5'],
        ['Swimming (moderate)', '7.0'],
        ['Walking (brisk)', '4.3'],
        ['Jump Rope', '12.3'],
        ['Elliptical', '5.0'],
    ]
)
doc.add_paragraph('Formula: Calories = MET x weight_kg x duration_hours')

add_heading_styled('5.8 Custom Daily Targets', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Create Target', 'Name (text) + optional description'],
        ['Examples', '"Read 30 minutes", "Meditate 10 min", "No sugar", "Take vitamins", "Stretch"'],
        ['Toggle', 'Simple checkbox - done or not done'],
        ['Max Targets', '10 custom targets per day'],
        ['Persistence', 'Targets repeat daily until user deletes them'],
        ['Completion', 'Included in daily completion percentage calculation'],
    ]
)

# ============================================================
# 6. PROGRESS & ANALYTICS
# ============================================================
add_heading_styled('6. Progress & Analytics', level=1)

add_heading_styled('6.1 Charts Overview', level=2)
add_table_from_data(
    ['Chart', 'Type', 'Data', 'Timeframe Options'],
    [
        ['Weight Trend', 'Line chart', 'Daily weight entries', '1W, 1M, 3M, 6M, All'],
        ['Calorie Trend', 'Bar chart', 'Daily calories vs target (target as horizontal line)', '1W, 1M'],
        ['Macro Breakdown', 'Donut/Pie chart', 'Average protein/carbs/fats over selected period', '1W, 1M'],
        ['Weekly Completion', 'Bar chart', 'Daily completion % for each day of week', 'Current week, last 4 weeks'],
        ['Workout Volume', 'Line chart', 'Total volume per workout session over time', '1M, 3M'],
        ['Sleep Trend', 'Bar chart (colored)', 'Hours per night, colored by quality rating', '1W, 1M'],
        ['Water Consistency', 'Heatmap calendar', 'Days where water goal met (green) vs missed (gray)', 'Current month'],
        ['Steps Trend', 'Bar chart', 'Daily steps vs goal', '1W, 1M'],
    ]
)

add_heading_styled('6.2 Best & Worst Habit Analysis', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Calculation', 'Completion rate per tracker over last 30 days'],
        ['Best Habit', 'Tracker with highest completion rate. E.g., "Water Intake - 93% consistency!"'],
        ['Worst Habit', 'Tracker with lowest completion rate. E.g., "Sleep - only 45% on target"'],
        ['Display', 'Card with trophy (best) and target (worst) icons, with improvement suggestions'],
        ['Suggestion', 'Pre-written tips per tracker. E.g., worst=sleep -> "Try setting a bedtime alarm"'],
    ]
)

add_heading_styled('6.3 Weekly Summary', level=2)
doc.add_paragraph('Generated every Sunday (or viewable anytime for the current week).')
add_table_from_data(
    ['Metric', 'Display'],
    [
        ['Average Daily Calories', '"1,850 avg / 1,900 target"'],
        ['Average Weight', '"74.3 kg (-0.5 from last week)"'],
        ['Total Workouts', '"5 workouts this week"'],
        ['Avg Sleep', '"7.2 hrs average"'],
        ['Water Consistency', '"6/7 days goal met"'],
        ['Step Average', '"8,200 avg steps"'],
        ['Streak', '"Current streak: 12 days"'],
        ['Completion Rate', '"Week average: 82%"'],
        ['Best Day', '"Tuesday - 100% completion!"'],
    ]
)

# ============================================================
# 7. STREAKS & MOTIVATION
# ============================================================
add_heading_styled('7. Streaks & Motivation', level=1)

add_heading_styled('7.1 Streak System', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Display', 'Fire emoji + number + "day streak"'],
        ['Location', 'Dashboard top-right, and Progress page header'],
        ['Animation', 'Flame grows/pulses at milestones (7, 14, 30, 60, 90, 100, 365)'],
        ['Streak Freeze', 'Post-MVP: 1 free freeze per week'],
    ]
)

add_heading_styled('7.2 Achievement Badges', level=2)
add_table_from_data(
    ['Badge', 'Condition', 'Icon'],
    [
        ['First Step', 'Complete onboarding', 'Flag'],
        ['Week Warrior', '7-day streak', 'Fire'],
        ['Fortnight Fighter', '14-day streak', 'Swords'],
        ['Month Master', '30-day streak', 'Trophy'],
        ['Century Club', '100-day streak', '100'],
        ['Hydration Hero', '30 consecutive days water goal met', 'Water drop'],
        ['Early Bird', 'Log sleep before 8 AM for 7 days', 'Bird'],
        ['Iron Will', '20 gym workouts in a month', 'Flexed bicep'],
        ['Marathon Mind', 'Track for 365 days total (not consecutive)', 'Medal'],
        ['Macro Master', 'Hit macro targets (+/-5%) for 7 consecutive days', 'Target'],
        ['PR Crusher', 'Set 10 personal records', 'Weight lifter'],
        ['Night Owl Cured', 'Sleep >=7hrs for 14 consecutive days', 'Sleep'],
    ]
)

add_heading_styled('7.3 Motivational Quotes', level=2)
add_table_from_data(
    ['Feature', 'Specification'],
    [
        ['Source', 'Built-in array of 100+ fitness/health quotes stored in database seed'],
        ['Rotation', 'One per day, selected by dayOfYear % quotes.length'],
        ['Display', 'Dashboard greeting banner, italic text with author'],
        ['API Endpoint', 'GET /api/quote returns today\'s quote'],
    ]
)

# ============================================================
# 8. UI LAYOUT & NAVIGATION
# ============================================================
add_heading_styled('8. UI Layout & Navigation', level=1)

add_heading_styled('8.1 Overall Structure', level=2)
doc.add_paragraph('The application is a single-page React app with bottom tab navigation (mobile) / sidebar navigation (desktop).')
add_bullet('Mobile: Bottom nav bar with 4 tabs + center FAB')
add_bullet('Desktop: Left sidebar nav + main content area')
add_bullet('All transitions powered by Framer Motion (slide/fade)')
add_bullet('Tailwind CSS for responsive design')
doc.add_paragraph()
add_table_from_data(
    ['Spec', 'Value'],
    [
        ['Minimum resolution', '360x640 (mobile-first), scales up to desktop'],
        ['Font', 'Inter (Google Fonts)'],
        ['Theme', 'Dark mode (default) / Light mode toggle via Tailwind'],
    ]
)

doc.add_paragraph('Color Palette:')
add_table_from_data(
    ['Element', 'Dark Mode', 'Light Mode'],
    [
        ['Background', '#0F0F0F', '#F5F5F5'],
        ['Card Background', '#1A1A2E', '#FFFFFF'],
        ['Primary Accent', '#00D4AA (teal/green)', '#00B894'],
        ['Secondary Accent', '#6C5CE7 (purple)', '#6C5CE7'],
        ['Danger/Over', '#FF6B6B', '#E74C3C'],
        ['Warning', '#FDCB6E', '#F39C12'],
        ['Success', '#00D4AA', '#00B894'],
        ['Text Primary', '#FFFFFF', '#2D3436'],
        ['Text Secondary', '#A0A0B0', '#636E72'],
    ]
)

add_heading_styled('8.2 Navigation Structure', level=2)
add_table_from_data(
    ['Tab', 'Icon', 'Label', 'Content'],
    [
        ['1', 'Home', 'Home', 'Daily Dashboard'],
        ['2', 'Pencil', 'Log', 'All daily trackers (tabbed: Food / Workout / Cardio)'],
        ['3', 'Plus', '-', 'FAB: Quick add (centered, raised)'],
        ['4', 'Chart', 'Progress', 'Charts & analytics'],
        ['5', 'User', 'Profile', 'Settings, achievements, profile edit'],
    ]
)

add_heading_styled('8.3 Page Breakdown', level=2)
add_table_from_data(
    ['Page', 'Route', 'Content'],
    [
        ['Onboarding', '/onboarding', 'Multi-step questionnaire (only shown once)'],
        ['Dashboard', '/ (home)', 'All daily overview widgets'],
        ['Food Logger', '/log/food', 'Meal categories + food entry form'],
        ['Workout Logger', '/log/workout', 'Exercise builder + set tracker'],
        ['Cardio Logger', '/log/cardio', 'Activity type + duration/distance'],
        ['Progress', '/progress', 'All charts + weekly summary + best/worst habit'],
        ['Achievements', '/achievements', 'Badge grid + streak history'],
        ['Profile', '/profile', 'Edit profile, change goals, recalculate TDEE'],
        ['Settings', '/settings', 'Units, theme, targets, macro split, data management'],
    ]
)

# ============================================================
# 9. SETTINGS & PROFILE
# ============================================================
add_heading_styled('9. Settings & Profile', level=1)

add_heading_styled('9.1 Profile Settings', level=2)
add_table_from_data(
    ['Setting', 'Options', 'Default'],
    [
        ['Name', 'Text input', 'From onboarding'],
        ['Age', 'Number', 'From onboarding'],
        ['Gender', 'Male / Female / Other', 'From onboarding'],
        ['Weight', 'Number + unit toggle (kg/lbs)', 'From onboarding'],
        ['Height', 'Number + unit toggle (cm or ft/in)', 'From onboarding'],
        ['Activity Level', '5-level dropdown', 'From onboarding'],
        ['Fitness Goal', 'Cut / Bulk / Maintain', 'From onboarding'],
    ]
)
add_note('On any change: TDEE recalculated automatically via backend. Calorie and macro targets update.')

add_heading_styled('9.2 Target Settings', level=2)
add_table_from_data(
    ['Setting', 'Options', 'Default'],
    [
        ['Daily Calories', 'Auto (from TDEE) or Manual override', 'Auto'],
        ['Macro Split', '3 sliders (must total 100%)', 'Based on goal'],
        ['Water Goal', 'Glasses or Liters', '8 glasses (2L)'],
        ['Step Goal', 'Number', '10,000'],
        ['Sleep Goal', 'Hours', '8'],
        ['Streak Threshold', 'Completion % required', '80%'],
    ]
)

add_heading_styled('9.3 App Settings', level=2)
add_table_from_data(
    ['Setting', 'Options', 'Default'],
    [
        ['Theme', 'Dark / Light / System', 'Dark'],
        ['Units (Weight)', 'kg / lbs', 'kg'],
        ['Units (Height)', 'cm / ft-in', 'cm'],
        ['Units (Distance)', 'km / mi', 'km'],
        ['Date Format', 'DD/MM/YYYY or MM/DD/YYYY', 'DD/MM/YYYY'],
        ['Reset All Data', 'Confirmation dialog -> wipe database', '-'],
        ['Re-do Onboarding', 'Restart questionnaire', '-'],
        ['Export Data', 'Download as JSON file', '-'],
    ]
)

# ============================================================
# 10. TECHNICAL ARCHITECTURE
# ============================================================
doc.add_page_break()
add_heading_styled('10. Technical Architecture', level=1)
add_new_section_tag()

add_heading_styled('10.1 Tech Stack', level=2)

doc.add_paragraph('Frontend:')
add_table_from_data(
    ['Technology', 'Purpose', 'Why'],
    [
        ['React 18 + TypeScript', 'UI components', 'Industry standard, component-based, type-safe'],
        ['Vite', 'Build tool & dev server', 'Starts in <1 second, fast hot module replacement'],
        ['Tailwind CSS', 'Styling', 'Utility-first CSS, built-in dark mode, 3-5x faster than writing CSS files'],
        ['Framer Motion', 'Animations & transitions', 'Best React animation library, simple API, premium feel'],
        ['Zustand', 'State management', 'Simpler than Redux, minimal boilerplate, used in Grandline'],
        ['Recharts', 'Progress charts', 'Built for React, beautiful defaults, easy to customize'],
        ['Lucide React', 'Icons', '1000+ clean icons, lightweight, consistent style'],
        ['Axios', 'HTTP requests to backend', 'Better error handling than fetch, request/response interceptors'],
    ]
)

doc.add_paragraph('Backend:')
add_table_from_data(
    ['Technology', 'Purpose', 'Why'],
    [
        ['Node.js + Express', 'REST API server', 'Same language as frontend (JavaScript/TypeScript everywhere)'],
        ['TypeScript', 'Type safety', 'Catches bugs before runtime, better IDE support'],
        ['PostgreSQL', 'Database', 'Rock-solid relational DB, handles complex queries (aggregations, trends)'],
        ['Prisma', 'ORM', 'Auto-generated types, easy migrations, Prisma Studio for debugging'],
        ['Zod', 'Request validation', 'Runtime validation with TypeScript type inference'],
        ['bcrypt + JWT', 'Authentication', 'Future-ready for multi-user auth (not used in MVP)'],
        ['cors', 'Cross-origin requests', 'Allows frontend (port 5173) to talk to backend (port 3000)'],
        ['dotenv', 'Environment variables', 'Keeps database URL and secrets out of code'],
    ]
)

doc.add_paragraph('Dev & Testing Tools:')
add_table_from_data(
    ['Tool', 'Purpose'],
    [
        ['Vitest', 'Unit & integration tests (fast, works with TypeScript)'],
        ['Supertest', 'API endpoint testing (simulates HTTP requests)'],
        ['tsx', 'Run TypeScript directly without compiling'],
        ['nodemon', 'Auto-restart server on file changes during development'],
        ['ESLint + Prettier', 'Code quality and formatting'],
    ]
)

add_heading_styled('10.2 System Architecture', level=2)
add_code_block(
    'YOUR BROWSER (localhost:5173)\n'
    '  React + Tailwind + Framer Motion -> What you see\n'
    '  Zustand -> Remembers data in-app\n'
    '  Recharts -> Draws charts\n'
    '  Axios -> Sends requests to server\n'
    '         |\n'
    '         | HTTP (REST API)\n'
    '         v\n'
    'YOUR COMPUTER (localhost:3000)\n'
    '  Express -> Receives requests\n'
    '  Zod -> Validates data\n'
    '  Prisma -> Talks to database\n'
    '         |\n'
    '         | SQL queries\n'
    '         v\n'
    'PostgreSQL (localhost:5432)\n'
    '  Stores everything permanently\n'
)

add_heading_styled('10.3 How They Connect (Example: Logging a Workout)', level=2)
add_table_from_data(
    ['Step', 'Who', 'Does What'],
    [
        ['1', 'You', 'Fill out "Bench Press, 80kg, 3 sets" and click Save'],
        ['2', 'React (frontend)', 'Packages the data and sends POST /api/workouts to server'],
        ['3', 'Express (server)', 'Receives the request'],
        ['4', 'Zod (validator)', 'Checks: is weight a number? is sets > 0? Rejects bad data.'],
        ['5', 'Prisma (ORM)', 'Converts to SQL and saves to PostgreSQL database'],
        ['6', 'PostgreSQL', 'Stores the workout permanently on disk'],
        ['7', 'Express', 'Sends back { success: true, workout: {...} }'],
        ['8', 'Zustand (state)', 'Updates the app\'s memory -> React redraws the dashboard'],
    ]
)

add_heading_styled('10.4 Project Structure', level=2)
add_code_block(
    'Project-Hercules/\n'
    '  server/                      - Backend\n'
    '    prisma/\n'
    '      schema.prisma            - Database schema (all tables)\n'
    '      seed.ts                  - Sample data (quotes, exercises)\n'
    '      migrations/              - Auto-generated by Prisma\n'
    '    src/\n'
    '      index.ts                 - Express app entry point\n'
    '      config/env.ts            - Environment config\n'
    '      routes/                  - API route definitions\n'
    '        profile.routes.ts\n'
    '        meals.routes.ts\n'
    '        workouts.routes.ts\n'
    '        cardio.routes.ts\n'
    '        water.routes.ts\n'
    '        sleep.routes.ts\n'
    '        steps.routes.ts\n'
    '        targets.routes.ts\n'
    '        progress.routes.ts\n'
    '        settings.routes.ts\n'
    '      controllers/             - Route handlers\n'
    '      services/                - Business logic (TDEE, streaks, badges)\n'
    '      validators/              - Zod schemas for validation\n'
    '      middleware/              - Error handler, validation middleware\n'
    '      utils/                   - TDEE calc, streak logic, MET values\n'
    '      types/                   - Shared TypeScript types\n'
    '    tests/\n'
    '      routes/                  - API endpoint tests\n'
    '      services/                - Business logic tests\n'
    '      setup.ts                 - Test database setup\n'
    '    package.json\n'
    '    tsconfig.json\n'
    '    .env\n'
    '\n'
    '  client/                      - Frontend\n'
    '    src/\n'
    '      main.tsx                 - React entry point\n'
    '      App.tsx                  - Router + layout\n'
    '      api/                     - Axios API client functions\n'
    '      components/\n'
    '        ui/                    - Buttons, cards, inputs, modals\n'
    '        dashboard/             - Dashboard widgets\n'
    '        trackers/              - Food, workout, water, sleep, steps\n'
    '        progress/              - Charts & analytics\n'
    '        onboarding/            - Questionnaire steps\n'
    '      pages/                   - Page components\n'
    '      stores/                  - Zustand stores\n'
    '      hooks/                   - Custom React hooks\n'
    '      utils/                   - Helpers (date, units)\n'
    '      types/                   - TypeScript types\n'
    '    tailwind.config.ts\n'
    '    vite.config.ts\n'
    '    package.json\n'
)

# ============================================================
# 11. DATABASE SCHEMA
# ============================================================
add_heading_styled('11. Database Schema (Prisma)', level=1)
add_new_section_tag()

doc.add_paragraph('The database contains 14 tables, managed via Prisma ORM:')
add_table_from_data(
    ['Table', 'Purpose', 'Key Fields'],
    [
        ['Profile', 'User profile + TDEE data', 'name, age, gender, weight, height, activityLevel, goal, bmr, tdee, calorieTarget, macros'],
        ['Meal', 'Food entries per day', 'date, category (breakfast/lunch/dinner/snacks), name, calories, protein, carbs, fats, isFavorite'],
        ['Workout', 'Gym workout sessions', 'date, name, duration, notes, isTemplate'],
        ['Exercise', 'Exercises within a workout', 'name, category (Chest/Back/etc.), notes, sortOrder'],
        ['ExerciseSet', 'Sets within an exercise', 'reps, weight, type (normal/warmup/dropset/failure)'],
        ['Cardio', 'Cardio sessions', 'date, activityType, duration, distance, caloriesBurned'],
        ['WaterLog', 'Daily water intake', 'date, amount, unit (unique per profile+date)'],
        ['SleepLog', 'Daily sleep tracking', 'date, hours, quality, bedtime, wakeTime (unique per profile+date)'],
        ['StepLog', 'Daily step count', 'date, count (unique per profile+date)'],
        ['WeightLog', 'Daily weight entry', 'date, weight (unique per profile+date)'],
        ['CustomTarget', 'User-defined daily targets', 'name, isActive'],
        ['DailyLog', 'Daily completion tracking', 'date, completionPct, targetsMet (JSON)'],
        ['Achievement', 'Earned badges', 'badge (e.g., "week_warrior"), earnedAt'],
        ['Settings', 'App preferences', 'theme, units, goals, streakThreshold'],
    ]
)

doc.add_paragraph('Key relationships:')
add_bullet('Profile -> has many: Meals, Workouts, Cardio, WaterLogs, SleepLogs, StepLogs, WeightLogs, CustomTargets, DailyLogs, Achievements')
add_bullet('Profile -> has one: Settings')
add_bullet('Workout -> has many: Exercises (cascade delete)')
add_bullet('Exercise -> has many: ExerciseSets (cascade delete)')
add_bullet('WaterLog, SleepLog, StepLog, WeightLog, DailyLog: unique constraint on [profileId, date] (one entry per day)')

# ============================================================
# 12. API ENDPOINTS
# ============================================================
add_heading_styled('12. API Endpoints', level=1)
add_new_section_tag()

add_heading_styled('12.1 Profile & Onboarding', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/profile', 'Create profile (onboarding complete, calculates TDEE)'],
        ['GET', '/api/profile/:id', 'Get profile with TDEE and macro data'],
        ['PUT', '/api/profile/:id', 'Update profile (recalculates TDEE automatically)'],
    ]
)

add_heading_styled('12.2 Meals', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/meals', 'Log a meal'],
        ['GET', '/api/meals?date=YYYY-MM-DD', 'Get meals for a date'],
        ['PUT', '/api/meals/:id', 'Edit a meal'],
        ['DELETE', '/api/meals/:id', 'Delete a meal'],
        ['GET', '/api/meals/favorites', 'Get favorite foods'],
        ['GET', '/api/meals/recent', 'Get recent foods (last 20)'],
    ]
)

add_heading_styled('12.3 Workouts', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/workouts', 'Log a workout (with nested exercises + sets)'],
        ['GET', '/api/workouts?date=YYYY-MM-DD', 'Get workouts for a date'],
        ['GET', '/api/workouts/:id', 'Get workout detail with exercises and sets'],
        ['DELETE', '/api/workouts/:id', 'Delete a workout (cascades to exercises + sets)'],
        ['GET', '/api/workouts/templates', 'Get saved workout templates'],
        ['GET', '/api/workouts/records', 'Get personal records per exercise'],
    ]
)

add_heading_styled('12.4 Cardio', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/cardio', 'Log a cardio session (auto-calculates calories via MET if not provided)'],
        ['GET', '/api/cardio?date=YYYY-MM-DD', 'Get cardio sessions for a date'],
        ['DELETE', '/api/cardio/:id', 'Delete a cardio session'],
    ]
)

add_heading_styled('12.5 Daily Trackers (Water, Sleep, Steps, Weight)', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/water', 'Log/update water for a date (upsert)'],
        ['GET', '/api/water?date=YYYY-MM-DD', 'Get water for a date'],
        ['POST', '/api/sleep', 'Log/update sleep for a date (upsert)'],
        ['GET', '/api/sleep?date=YYYY-MM-DD', 'Get sleep for a date'],
        ['POST', '/api/steps', 'Log/update steps for a date (upsert)'],
        ['GET', '/api/steps?date=YYYY-MM-DD', 'Get steps for a date'],
        ['POST', '/api/weight', 'Log weight for a date (upsert)'],
        ['GET', '/api/weight?range=1M', 'Get weight history for charting'],
    ]
)

add_heading_styled('12.6 Custom Targets', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['POST', '/api/targets', 'Create a custom target'],
        ['GET', '/api/targets', 'Get all active targets'],
        ['PUT', '/api/targets/:id', 'Update target (name or active status)'],
        ['DELETE', '/api/targets/:id', 'Delete target'],
        ['POST', '/api/targets/daily', 'Log daily target completions (JSON map)'],
    ]
)

add_heading_styled('12.7 Dashboard & Progress', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['GET', '/api/dashboard?date=YYYY-MM-DD', 'Get ALL data for a single day (meals, water, sleep, steps, workouts, cardio, targets, completion, streak, quote)'],
        ['GET', '/api/progress/weight?range=1M', 'Weight trend data for chart'],
        ['GET', '/api/progress/calories?range=1W', 'Daily calorie totals for chart'],
        ['GET', '/api/progress/macros?range=1W', 'Average macro breakdown for chart'],
        ['GET', '/api/progress/completion?range=1W', 'Daily completion percentages for chart'],
        ['GET', '/api/progress/habits', 'Best/worst habit analysis (30-day window)'],
        ['GET', '/api/progress/summary', 'Weekly summary (all metrics aggregated)'],
    ]
)

add_heading_styled('12.8 Streaks, Achievements, Settings, Quotes', level=2)
add_table_from_data(
    ['Method', 'Endpoint', 'Purpose'],
    [
        ['GET', '/api/streaks', 'Get current streak + best streak'],
        ['GET', '/api/achievements', 'Get all badges (earned + locked with conditions)'],
        ['GET', '/api/settings', 'Get app settings'],
        ['PUT', '/api/settings', 'Update app settings'],
        ['GET', '/api/quote', 'Get today\'s motivational quote'],
    ]
)

# ============================================================
# 13. EDGE CASES
# ============================================================
add_heading_styled('13. Edge Cases', level=1)
add_table_from_data(
    ['Scenario', 'Expected Behavior', 'Error Handling'],
    [
        ['User enters 0 or negative weight/height', 'Zod validation rejects request (400)', 'Inline red text: "Please enter a valid number"'],
        ['Macro sliders don\'t total 100%', 'Save button disabled until 100%', 'Live counter: "Total: 95% (must be 100%)"'],
        ['User logs >10,000 calories', 'Allowed but confirmation dialog', '"That\'s a lot! Are you sure?"'],
        ['Database connection lost', 'API returns 500 with friendly message', 'Toast: "Server error. Please try again."'],
        ['User hasn\'t logged in 30+ days', 'Streak resets. Welcome-back message.', '"Welcome back! Let\'s get back on track."'],
        ['User changes goal from Cut to Bulk', 'TDEE recalculated server-side', 'Confirmation: "Your target will change from X to Y"'],
        ['User changes weight unit mid-use', 'All stored weights converted', 'Conversion happens on unit toggle'],
        ['No data for a chart', 'API returns empty array', 'Empty state: "No data yet. Start tracking!"'],
        ['Date rolls over at midnight', 'New day auto-detected on dashboard load', 'Fresh data loaded from backend'],
        ['Multiple workouts in one day', 'All saved as separate entries', 'Dashboard shows total count'],
        ['Invalid API request (bad JSON)', 'Zod rejects with 400 + field-level errors', 'Frontend displays specific field errors'],
        ['Concurrent writes to same daily log', 'Last write wins (upsert)', 'No conflict - single user MVP'],
    ]
)

# ============================================================
# 14. LOADING, EMPTY & ERROR STATES
# ============================================================
add_heading_styled('14. Loading, Empty & Error States', level=1)

add_heading_styled('14.1 Loading States', level=2)
add_table_from_data(
    ['View', 'Loading UX'],
    [
        ['App startup', 'Hercules logo with pulsing animation (1-2 seconds max)'],
        ['Dashboard widgets', 'Skeleton cards with Tailwind shimmer effect'],
        ['Charts', 'Spinner inside chart container + "Crunching numbers..."'],
        ['Saving data', 'Brief toast: "Saved!" (auto-dismiss 1.5s)'],
        ['API requests', 'Loading state in Zustand store -> skeleton UI'],
    ]
)

add_heading_styled('14.2 Empty States', level=2)
add_table_from_data(
    ['View', 'Condition', 'Empty State UX'],
    [
        ['Dashboard', 'No data logged today', '"Your day starts here! Tap + to log your first meal or workout."'],
        ['Food log', 'No meals logged', 'Plate icon + "No meals logged yet. What did you eat today?"'],
        ['Workout log', 'No workouts logged', 'Dumbbell icon + "No workouts yet. Time to hit the gym?"'],
        ['Progress charts', 'Less than 3 days of data', '"Track for a few more days to see your trends."'],
        ['Achievements', 'No badges earned', 'All badges shown grayed out with conditions'],
        ['Weekly summary', 'Less than 1 full week', '"Complete your first full week to see your summary!"'],
    ]
)

add_heading_styled('14.3 Error States', level=2)
add_table_from_data(
    ['Error Type', 'Trigger', 'User Message', 'Recovery'],
    [
        ['API Failure (500)', 'Server error', '"Something went wrong. Please try again."', 'Retry button + auto-retry after 5s'],
        ['Network Offline', 'No connectivity', 'Top banner: "You are offline."', 'Auto-reconnect when online'],
        ['Validation Error (400)', 'Bad input data', 'Inline field-level error messages', 'Fix and retry'],
        ['Not Found (404)', 'Invalid profile/resource ID', '"Resource not found."', 'Redirect to dashboard'],
    ]
)

# ============================================================
# 15. ACCESSIBILITY
# ============================================================
add_heading_styled('15. Accessibility', level=1)
add_table_from_data(
    ['Concern', 'Approach'],
    [
        ['Color contrast', 'WCAG AA minimum (4.5:1 for text)'],
        ['Color-blind safety', 'Never rely on color alone - always include icons/text labels'],
        ['Keyboard navigation', 'All interactive elements focusable via Tab. Enter/Space to activate.'],
        ['Screen reader', 'Semantic HTML, ARIA labels on icons and charts'],
        ['Touch targets', 'Minimum 44x44px for all tap targets (mobile)'],
        ['Font size', 'Base 16px, scalable. Never smaller than 12px.'],
        ['Reduced motion', 'Framer Motion respects prefers-reduced-motion'],
    ]
)

# ============================================================
# 16. IMPLEMENTATION ROADMAP
# ============================================================
doc.add_page_break()
add_heading_styled('16. Implementation Roadmap', level=1)
add_new_section_tag()

doc.add_paragraph('Build order: Database -> Backend -> Backend Testing -> Frontend -> Integration Testing')
doc.add_paragraph()

phases = [
    ('Phase 1 - Database Setup', [
        'Initialize server/ project (npm init, install dependencies)',
        'Configure Prisma + PostgreSQL connection (.env)',
        'Write schema.prisma with all 14 models',
        'Run prisma migrate dev to create all tables',
        'Write seed.ts with sample data (quotes, exercise categories, sample profile)',
        'Verify database with Prisma Studio',
    ]),
    ('Phase 2 - Backend Core API', [
        'Express app setup (index.ts, middleware, error handler, CORS)',
        'Profile routes + controller + service (TDEE calculation)',
        'Meals routes + controller + service (CRUD + favorites + recent)',
        'Workouts routes + controller + service (nested exercises + sets)',
        'Cardio routes + controller + service (MET calorie estimation)',
        'Water routes + controller + service (upsert per date)',
        'Sleep routes + controller + service (upsert per date)',
        'Steps routes + controller + service (upsert per date)',
        'Weight routes + controller + service (upsert per date)',
        'Custom targets routes + controller + service',
        'Zod validators for all endpoints',
    ]),
    ('Phase 3 - Backend Advanced API', [
        'Dashboard endpoint (aggregates ALL daily data in one call)',
        'Progress endpoints (weight trend, calorie trend, macros, completion)',
        'Streak calculation service (based on daily completion data)',
        'Achievement/badge detection service (checked on each log save)',
        'Best/worst habit analysis (30-day completion rates per tracker)',
        'Weekly summary endpoint (all metrics aggregated)',
        'Settings endpoints (GET/PUT)',
        'Quote of the day endpoint (from seeded quotes)',
    ]),
    ('Phase 4 - Backend Testing', [
        'Test setup (test database, cleanup between tests)',
        'TDEE calculation unit tests (all formulas, all activity levels)',
        'Streak logic unit tests (increment, reset, best streak)',
        'Achievement detection unit tests (all 12 badges)',
        'MET calorie estimation unit tests (all activity types)',
        'Profile API tests (create, read, update, TDEE recalculation)',
        'Meals API tests (CRUD + favorites + recent)',
        'Workouts API tests (nested create with exercises + sets, cascade delete)',
        'All other route tests (water, sleep, steps, cardio, weight, targets)',
        'Dashboard aggregation test (verify all data combined correctly)',
        'Progress endpoint tests (verify chart data shapes)',
        'Edge case tests (invalid data, missing fields, boundary values)',
    ]),
    ('Phase 5 - Frontend Foundation', [
        'Initialize client/ project (Vite + React + TypeScript)',
        'Tailwind CSS setup + theme config (dark/light, color palette from Section 8)',
        'App layout (bottom nav mobile, sidebar desktop)',
        'React Router setup (all routes from Section 8.3)',
        'Zustand stores (profile, daily data, settings)',
        'Axios API client (base URL config + all endpoint functions)',
        'Reusable UI components (Button, Card, Input, Modal, ProgressRing, Toast)',
        'Framer Motion page transition wrapper',
    ]),
    ('Phase 6 - Frontend: Onboarding', [
        'Welcome screen with Framer Motion animation',
        'Multi-step questionnaire (name, body stats, activity level, goal)',
        'Results screen (TDEE display, calorie/macro breakdown)',
        'Progress dots indicator between steps',
        'Zod-based input validation per step',
        'POST to /api/profile on completion -> redirect to dashboard',
    ]),
    ('Phase 7 - Frontend: Dashboard & Trackers', [
        'Dashboard page with all 11 widgets from Section 5.1',
        'Calorie ring (SVG circular progress with Framer Motion)',
        'Macro bars (3 horizontal Tailwind progress bars)',
        'Water tracker (glass grid with fill animation)',
        'Sleep card, Steps card (circular arc), Workout summary card',
        'Custom targets checklist with checkboxes',
        'Streak counter with fire animation at milestones',
        'Daily completion bar (colored by percentage)',
        'FAB (floating action button) with radial expand',
        'Food logger (meal category tabs, entry form, recent, favorites)',
        'Workout logger (exercise builder, sets, rest timer, templates)',
        'Cardio logger (activity type dropdown, duration, distance)',
    ]),
    ('Phase 8 - Frontend: Progress & Polish', [
        'Weight trend chart (Recharts line chart)',
        'Calorie trend chart (Recharts bar chart with target line)',
        'Macro breakdown chart (Recharts donut)',
        'Weekly completion chart (Recharts bar)',
        'Sleep trend chart (colored bars by quality)',
        'Best/worst habit cards with improvement suggestions',
        'Weekly summary view with all metrics',
        'Achievement badge grid (earned = color, locked = gray)',
        'Profile edit page (recalculates TDEE on save)',
        'Settings page (theme toggle, units, goals, export, reset)',
        'Empty states for all views (from Section 14.2)',
        'Loading skeletons (Tailwind shimmer)',
        'Error handling + toast notifications',
        'Responsive design pass (360px to desktop)',
        'Framer Motion animation polish (page transitions, card enters, celebrations)',
    ]),
    ('Phase 9 - Integration Testing', [
        'Full user journey: onboarding -> dashboard -> log data -> view progress',
        'Cross-browser testing (Chrome, Edge, Firefox)',
        'Mobile responsiveness testing (360px, 390px, 768px, 1280px)',
        'Edge case testing (all scenarios from Section 13)',
        'Performance check (API response times, chart rendering)',
    ]),
]

for phase_title, items in phases:
    add_heading_styled(phase_title, level=2)
    for item in items:
        add_bullet(item)

# ============================================================
# 17. VERIFICATION
# ============================================================
add_heading_styled('17. Verification & Testing', level=1)

add_heading_styled('17.1 How to Run Locally', level=2)
add_code_block(
    'Terminal 1: PostgreSQL must be running\n\n'
    'Terminal 2: Backend\n'
    '  cd server\n'
    '  npm install\n'
    '  npx prisma migrate dev\n'
    '  npx prisma db seed\n'
    '  npm run dev\n'
    '  -> http://localhost:3000\n\n'
    'Terminal 3: Frontend\n'
    '  cd client\n'
    '  npm install\n'
    '  npm run dev\n'
    '  -> http://localhost:5173\n'
)

add_heading_styled('17.2 Phase Verification Checkpoints', level=2)
add_table_from_data(
    ['Phase', 'Verification'],
    [
        ['Phase 1', 'Open Prisma Studio -> all 14 tables exist with correct columns and relations'],
        ['Phase 2-3', 'Test all endpoints with curl/Postman -> correct responses'],
        ['Phase 4', 'Run npm test in server/ -> all tests pass, 0 failures'],
        ['Phase 5', 'Open localhost:5173 -> layout renders with navigation working'],
        ['Phase 6', 'Complete onboarding -> profile appears in database (check Prisma Studio)'],
        ['Phase 7', 'Log meals/workouts/water -> data persists, dashboard updates in real-time'],
        ['Phase 8', 'View charts with 3+ days of data -> all 8 charts render correctly'],
        ['Phase 9', 'Full journey works end-to-end without errors across browsers'],
    ]
)

add_heading_styled('17.3 Manual QA Checklist', level=2)
add_table_from_data(
    ['Test', 'Expected Result'],
    [
        ['Complete onboarding with valid data', 'Profile saved, TDEE calculated, redirected to dashboard'],
        ['Complete onboarding with invalid data', 'Validation prevents advancing (Zod errors shown)'],
        ['Log a meal and refresh page', 'Meal persists (stored in PostgreSQL)'],
        ['Log 5+ meals across all categories', 'Daily total correct, macros sum correctly'],
        ['Add 8 glasses of water', 'Visual shows all glasses filled, goal met'],
        ['Log sleep with bedtime/waketime', 'Hours auto-calculated correctly'],
        ['Log a gym workout with 3 exercises', 'Volume calculated, workout appears on dashboard'],
        ['Log a run of 5km in 30 min', 'Calories auto-estimated using MET formula (server-side)'],
        ['Create 3 custom targets, complete 2', 'Completion % reflects correctly'],
        ['Track for 7 consecutive days at >=80%', '"Week Warrior" badge earned'],
        ['Change weight unit from kg to lbs', 'All stored weights converted correctly'],
        ['Change goal from Cut to Bulk', 'Backend recalculates TDEE, frontend updates'],
        ['Export data', 'JSON file downloads with all data'],
        ['Reset all data', 'Database cleared, app returns to onboarding'],
        ['Open on mobile (360px)', 'Bottom nav visible, all content readable'],
        ['Toggle dark/light mode', 'All Tailwind theme classes switch correctly'],
        ['Kill backend server, try to log food', 'Error toast shown, no crash'],
    ]
)

# ============================================================
# APPENDIX A - CORE USER JOURNEYS
# ============================================================
doc.add_page_break()
add_heading_styled('Appendix A - Core User Journeys', level=1)

add_heading_styled('Journey 1: First-Time Setup', level=2)
for i, step in enumerate([
    'User opens localhost:5173 in Chrome',
    'React checks backend for profile -> none found -> redirect to /onboarding',
    'Welcome screen: "Welcome to Project Hercules!" with Framer Motion animation',
    'Clicks "Get Started"',
    'Enters name: "Rahul"',
    'Enters body stats: Age 25, Male, 75 kg, 175 cm',
    'Selects activity level: "Moderately Active"',
    'Selects goal: "Lose Weight (Cut)"',
    'Frontend POST /api/profile -> backend calculates BMR, TDEE, calorie target, macro split',
    'Results screen: "Your daily target: 2,171 calories" + macro split shown',
    'Clicks "Let\'s Go!" -> redirect to dashboard -> GET /api/dashboard loads empty day',
], 1):
    add_bullet(f'{i}. {step}')

add_heading_styled('Journey 2: Typical Daily Use', level=2)
for i, step in enumerate([
    'Opens app -> GET /api/dashboard?date=today -> "Good morning, Rahul!" + daily quote',
    'Taps sleep card -> logs 7.5 hrs, Good quality -> POST /api/sleep',
    'Taps FAB -> "Log Meal" -> enters food -> POST /api/meals -> dashboard calorie ring updates',
    'Taps FAB -> "Log Workout" -> creates "Push Day" -> POST /api/workouts (nested exercises + sets)',
    'Dashboard auto-refreshes: workout card + calorie ring updated',
    'Taps water glass icons throughout day -> POST /api/water (upsert)',
    'Logs lunch and dinner via food logger',
    'Before bed -> daily completion: 85% -> streak increments -> POST /api/targets/daily',
    'Optionally logs weight: POST /api/weight',
], 1):
    add_bullet(f'{i}. {step}')

add_heading_styled('Journey 3: Weekly Review', level=2)
for i, step in enumerate([
    'Opens Progress tab -> GET /api/progress/weight, /calories, /completion, /habits, /summary',
    'Weight trend chart: down 0.3 kg this week',
    'Calorie bars: mostly under target (good for cut)',
    'Macro donut: protein slightly low',
    'Weekly completion: 5/7 days at >=80%',
    'Best habit: Water (100%) -> worst: Sleep (57%)',
    'Achievement badge "Hydration Hero" newly earned!',
], 1):
    add_bullet(f'{i}. {step}')

# ============================================================
# APPENDIX B - UNIT CONVERSIONS
# ============================================================
add_heading_styled('Appendix B - Unit Conversion Reference', level=1)
add_table_from_data(
    ['Conversion', 'Formula'],
    [
        ['kg -> lbs', 'kg x 2.20462'],
        ['lbs -> kg', 'lbs x 0.453592'],
        ['cm -> ft/in', 'inches = cm / 2.54; ft = floor(inches/12); in = inches % 12'],
        ['ft/in -> cm', '(ft x 12 + in) x 2.54'],
        ['km -> mi', 'km x 0.621371'],
        ['mi -> km', 'mi x 1.60934'],
        ['glasses -> L', 'glasses x 0.25'],
        ['L -> glasses', 'L x 4'],
    ]
)

# ============================================================
# APPENDIX C - DESIGN INSPIRATION
# ============================================================
add_heading_styled('Appendix C - Design Inspiration', level=1)
add_table_from_data(
    ['Element', 'Direction'],
    [
        ['Overall vibe', 'Premium fitness app (MyFitnessPal meets Streaks meets Apple Health)'],
        ['Cards', 'Tailwind: rounded-xl, shadow-lg, backdrop-blur on dark mode'],
        ['Buttons', 'Pill-shaped (rounded-full), gradient accents on primary action'],
        ['Progress rings', 'SVG with Framer Motion animation, gradient stroke'],
        ['Typography', 'Inter - font-bold headings, font-normal body, font-light labels'],
        ['Spacing', 'Tailwind: p-4 to p-6, space-y-4 between cards'],
        ['Micro-interactions', 'Framer Motion: scale(0.95) on press, y-lift on hover, count-up numbers'],
        ['Celebration', 'Confetti animation on badge earn, streak milestone pulse'],
    ]
)

# ============================================================
# FOOTER
# ============================================================
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('End of Planning Document - Version 2.0 (Full-Stack)')
run.italic = True
run.font.size = Pt(10)

# Save
output_path = r'C:\Users\user\Desktop\Project-Hercules\Project_Hercules_Planning_Document_v2.0.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
